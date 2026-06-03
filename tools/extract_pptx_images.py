from pptx import Presentation
import openpyxl
import os

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX_PATH = os.path.join(BASE_DIR, 'picture-images.xlsx')
PPTX_PATH = os.path.join(BASE_DIR, 'picture-image-A5.pptx')
OUTPUT_DIR = os.path.join(BASE_DIR, 'images')

wb = openpyxl.load_workbook(XLSX_PATH)
ws = wb.active
names = [cell.value for cell in ws['B'][1:] if cell.value]

prs = Presentation(PPTX_PATH)
os.makedirs(OUTPUT_DIR, exist_ok=True)

saved = 0
skipped = 0
for i, slide in enumerate(prs.slides):
    if i >= len(names):
        break

    image_blob = None
    for rel in slide.part.rels.values():
        if 'image' in rel.reltype:
            image_blob = rel.target_part.blob
            break

    name = names[i]
    if image_blob is None:
        print(f'スライド{i+1} ({name}): 画像なし → スキップ')
        skipped += 1
        continue

    filepath = os.path.join(OUTPUT_DIR, f'{name}.png')
    with open(filepath, 'wb') as f:
        f.write(image_blob)
    print(f'保存: {filepath}')
    saved += 1

print(f'\n完了: {saved}枚保存, {skipped}枚スキップ')
